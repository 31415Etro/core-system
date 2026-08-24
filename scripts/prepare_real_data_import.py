from __future__ import annotations

import json
import re
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
PROJECT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT / "scripts" / "201_import_real_data.sql"
REPORT = PROJECT / "scripts" / "201_import_real_data_report.json"
DATA_DUMP = PROJECT / "scripts" / "201_import_real_data.json"
NAMESPACE = uuid.UUID("d59e3926-eb58-4e77-90b0-848df63e444f")
DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
MONEY_RE = re.compile(r"^-?[\d.]+,\d{2}$")


def uid(kind: str, key: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f"{kind}:{key}"))


def money(value: str) -> Decimal:
    return Decimal(value.replace(".", "").replace(",", "."))


def sql_text(value) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def iso_date(value: str) -> str:
    return datetime.strptime(value, "%d/%m/%Y").date().isoformat()


def extract_subscriptions():
    path = ROOT / "Assinaturas_Clientes_Consolidado.xlsx"
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["Assinaturas"]
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0] or str(row[0]).strip().upper() == "TOTAL":
            continue
        name, legal_name, document, scope, due_day, plan, plan_type, credits, ticket_types, monthly, payment, notes = row
        unlimited = str(credits).strip().lower() == "ilimitado"
        monthly_credits = None if unlimited else int(credits or 0)
        normalized_plan_type = "unlimited" if unlimited else "email" if "email" in str(plan_type).lower() else "server" if "servidor" in str(plan_type).lower() and monthly_credits == 0 else "credits"
        client_id = uid("client", str(document).strip())
        contract_id = uid("contract", str(document).strip())
        records.append({
            "client_id": client_id, "contract_id": contract_id, "name": str(name).strip(),
            "legal_name": str(legal_name).strip(), "document": str(document).strip(),
            "scope": str(scope or "").strip(), "due_day": int(due_day), "plan": str(plan).strip(),
            "plan_type": normalized_plan_type, "credits": monthly_credits, "unlimited": unlimited,
            "ticket_types": str(ticket_types or "").strip(), "monthly": Decimal(str(monthly)),
            "payment": str(payment or "").strip(), "notes": str(notes or "").strip(),
        })
    return records


def group_by_top(words):
    rows = {}
    for word in words:
        key = round(word["top"], 1)
        rows.setdefault(key, []).append(word)
    return rows


def extract_ailos(path: Path):
    transactions = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            words = page.extract_words(use_text_flow=False)
            rows = group_by_top(words)
            for top, row in sorted(rows.items()):
                date_words = [w for w in row if w["x0"] < 70 and DATE_RE.match(w["text"])]
                if not date_words:
                    continue
                date = date_words[0]["text"]
                desc = " ".join(w["text"] for w in sorted(row, key=lambda x: x["x0"]) if 70 <= w["x0"] < 290).strip()
                document = " ".join(w["text"] for w in sorted(row, key=lambda x: x["x0"]) if 290 <= w["x0"] < 370).strip()
                credit_words = [w["text"] for w in row if 370 <= w["x0"] < 455 and MONEY_RE.match(w["text"])]
                debit_words = [w["text"] for w in row if 455 <= w["x0"] < 530 and MONEY_RE.match(w["text"])]
                raw = credit_words[0] if credit_words else debit_words[0] if debit_words else None
                if raw is None:
                    continue
                amount = money(raw)
                sequence = len(transactions) + 1
                key = f"AILOS:{date}:{document}:{desc}:{amount}:{sequence}"
                transactions.append({"id": uid("transaction", key), "date": iso_date(date), "description": desc, "amount": amount.copy_abs(), "type": "income" if amount > 0 else "expense", "account": "AILOS", "external_id": document or f"page-{page_number}-{sequence}", "category": categorize(desc, amount, "AILOS")})
    return transactions


def extract_asaas(path: Path):
    transactions = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            words = page.extract_words(use_text_flow=False)
            dates = sorted([w for w in words if w["x0"] < 100 and DATE_RE.match(w["text"])], key=lambda w: w["top"])
            for index, date_word in enumerate(dates):
                top = date_word["top"]
                previous = dates[index - 1]["top"] if index else top - 24
                following = dates[index + 1]["top"] if index + 1 < len(dates) else top + 24
                lower = (previous + top) / 2
                upper = (top + following) / 2
                region = [w for w in words if lower <= w["top"] < upper]
                amount_words = [w["text"] for w in region if w["x0"] > 500 and MONEY_RE.match(w["text"])]
                if not amount_words:
                    continue
                amount = money(amount_words[-1])
                desc_words = [w for w in region if 130 <= w["x0"] < 500 and w["text"] not in ("R$",)]
                desc_words.sort(key=lambda w: (round(w["top"], 1), w["x0"]))
                desc = " ".join(w["text"] for w in desc_words).strip()
                invoice = re.search(r"fatura (?:nr\. )?(\d+)", desc, re.I)
                sequence = len(transactions) + 1
                external = invoice.group(1) if invoice else f"page-{page_number}-{sequence}"
                key = f"ASAAS:{date_word['text']}:{desc}:{amount}:{sequence}"
                transactions.append({"id": uid("transaction", key), "date": iso_date(date_word["text"]), "description": desc, "amount": amount.copy_abs(), "type": "income" if amount > 0 else "expense", "account": "ASAAS", "external_id": external, "category": categorize(desc, amount, "ASAAS")})
    return transactions


def categorize(description: str, amount: Decimal, account: str) -> str:
    text = description.casefold()
    if "waycore informatica" in text or "transação via pix com chave" in text:
        return "Transferência interna"
    if "cobrança recebida" in text or "ted recebida" in text:
        return "Receita de serviços"
    if "taxa" in text or "tar pct" in text:
        return "Taxas bancárias"
    if "estorno" in text or "devolucao" in text:
        return "Estorno"
    if "fatura cartao" in text:
        return "Cartão de crédito"
    if "darf" in text or "das" in text or "icms" in text or "secretaria de estado" in text:
        return "Impostos"
    if "celesc" in text or "unifique" in text:
        return "Utilidades"
    if amount > 0:
        return "Outras entradas"
    return "Outras saídas"


def normalize(value: str) -> str:
    import unicodedata
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def link_clients(transactions, subscriptions):
    aliases = {
        "ESCALA": ["escala rodrigues"], "VENEPAR": ["venepar"], "ELAI (BRX)": ["elai distribuidora"],
        "TATTICAS": ["tatticas publicidade"], "BIDUSKY": ["e r mendes", "er mendes"],
        "FARMA GRE": ["gretter"], "CLASS": ["class time"], "REBELLATO": ["dm farma"],
        "SMG": ["smg industrial"], "MAGNOREDES": ["magnoredes"], "ALPHA EXTERNO": ["alpha sul"],
        "GLOBAL LINK": ["global link"], "RODAN": ["erg transportes"], "ANDRAÇO": ["andraco comercio"],
        "TURISMAR": ["santiago cardoso"], "SHIPMAR": ["shipmar"],
        "NDI": ["ndi assessoria", "ndi importacao"], "OCEANIC": ["oceanic log"], "MOYO": ["moyo comercio"],
    }
    by_name = {item["name"]: item["client_id"] for item in subscriptions}
    for tx in transactions:
        text = normalize(tx["description"])
        tx["client_id"] = next((by_name[name] for name, patterns in aliases.items() if any(pattern in text for pattern in patterns)), None)


def build_sql(subscriptions, transactions, ailos_files, asaas_files):
    lines = ["-- CARGA REAL WAYCORE - gerada a partir dos arquivos fornecidos", "begin;", ""]
    for item in subscriptions:
        lines.append(
            "insert into public.clients (id,name,legal_name,document,status,notes) values "
            f"('{item['client_id']}',{sql_text(item['name'])},{sql_text(item['legal_name'])},{sql_text(item['document'])},'active',{sql_text(item['notes'])}) "
            "on conflict (id) do update set name=excluded.name,legal_name=excluded.legal_name,document=excluded.document,status=excluded.status,notes=excluded.notes,updated_at=now();"
        )
        scope_notes = " | ".join(filter(None, [item["scope"], f"Tipos: {item['ticket_types']}" if item["ticket_types"] else "", item["notes"]]))
        credits = "null" if item["credits"] is None else str(item["credits"])
        lines.append(
            "insert into public.support_contracts (id,client_id,plan_name,plan_type,monthly_credits,unlimited,monthly_value,due_day,payment_method,scope_notes,active) values "
            f"('{item['contract_id']}','{item['client_id']}',{sql_text(item['plan'])},{sql_text(item['plan_type'])},{credits},{str(item['unlimited']).lower()},{item['monthly']},{item['due_day']},{sql_text(item['payment'])},{sql_text(scope_notes)},true) "
            "on conflict (id) do update set plan_name=excluded.plan_name,plan_type=excluded.plan_type,monthly_credits=excluded.monthly_credits,unlimited=excluded.unlimited,monthly_value=excluded.monthly_value,due_day=excluded.due_day,payment_method=excluded.payment_method,scope_notes=excluded.scope_notes,active=true,updated_at=now();"
        )
    lines.append("")
    for tx in transactions:
        lines.append(
            "insert into public.financial_transactions (id,client_id,type,description,amount,transaction_date,paid_at,status,account_name,category,external_id) values "
            f"('{tx['id']}',{sql_text(tx['client_id'])},'{tx['type']}',{sql_text(tx['description'])},{tx['amount']},'{tx['date']}','{tx['date']} 12:00:00+00','paid',{sql_text(tx['account'])},{sql_text(tx['category'])},{sql_text(tx['external_id'])}) "
            "on conflict (id) do update set client_id=excluded.client_id,type=excluded.type,description=excluded.description,amount=excluded.amount,transaction_date=excluded.transaction_date,status='paid',account_name=excluded.account_name,category=excluded.category,external_id=excluded.external_id,updated_at=now();"
        )
    all_files = [(name, "AILOS", "saida", count) for name, count in ailos_files] + [(name, "ASAAS", "entrada", count) for name, count in asaas_files]
    file_names_sql = ",".join(sql_text(name) for name, *_ in all_files)
    lines.append(f"delete from public.bank_statement_imports where file_name in ({file_names_sql});")
    for name, account, direction, count in all_files:
        lines.append(
            f"insert into public.bank_statement_imports (account_name,direction,file_name,imported_rows,status) values "
            f"({sql_text(account)},{sql_text(direction)},{sql_text(name)},{count},'concluido');"
        )
    lines += [
        "", "commit;", "",
        "select 'clients' as item, count(*)::text as value from public.clients",
        "union all select 'contracts', count(*)::text from public.support_contracts",
        "union all select 'transactions', count(*)::text from public.financial_transactions",
        "union all select 'financial_net', coalesce(sum(case when type='income' then amount else -amount end),0)::text from public.financial_transactions;",
    ]
    OUTPUT.write_text("\n".join(lines), encoding="utf-8")


def main():
    subscriptions = extract_subscriptions()
    ailos_paths = sorted(set(ROOT.glob("Extrato Ailos*.pdf")) | set(ROOT.glob("extrato_conta_corrente*.pdf")))
    asaas_paths = sorted(ROOT.glob("Extrato Asaas*.pdf"))
    ailos_by_file = [(path.name, extract_ailos(path)) for path in ailos_paths]
    asaas_by_file = [(path.name, extract_asaas(path)) for path in asaas_paths]
    ailos = [tx for _, rows in ailos_by_file for tx in rows]
    asaas = [tx for _, rows in asaas_by_file for tx in rows]
    transactions = ailos + asaas
    link_clients(transactions, subscriptions)
    ailos_files = [(name, len(rows)) for name, rows in ailos_by_file]
    asaas_files = [(name, len(rows)) for name, rows in asaas_by_file]
    build_sql(subscriptions, transactions, ailos_files, asaas_files)
    def totals(rows):
        return {"rows": len(rows), "income": str(sum((x["amount"] for x in rows if x["type"] == "income"), Decimal())), "expense": str(sum((x["amount"] for x in rows if x["type"] == "expense"), Decimal())), "net": str(sum((x["amount"] if x["type"] == "income" else -x["amount"] for x in rows), Decimal()))}
    report = {"clients": len(subscriptions), "contracts": len(subscriptions), "monthly_contract_value": str(sum((x["monthly"] for x in subscriptions), Decimal())), "ailos": totals(ailos), "asaas": totals(asaas), "ailos_files": ailos_files, "asaas_files": asaas_files, "transactions_linked_to_clients": sum(1 for x in transactions if x.get("client_id")), "output": str(OUTPUT)}
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))

    def jsonify_sub(item):
        out = dict(item)
        out["monthly"] = str(out["monthly"])
        return out

    def jsonify_tx(tx):
        out = dict(tx)
        out["amount"] = str(out["amount"])
        return out

    dump = {
        "subscriptions": [jsonify_sub(s) for s in subscriptions],
        "transactions": [jsonify_tx(t) for t in transactions],
        "ailos_files": ailos_files, "asaas_files": asaas_files,
    }
    DATA_DUMP.write_text(json.dumps(dump, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
